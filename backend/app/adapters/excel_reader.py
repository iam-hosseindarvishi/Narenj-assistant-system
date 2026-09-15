"""Reads Excel workbooks (.xls via xlrd, .xlsx via openpyxl) and pasted TSV/CSV text
into row matrices compatible with the template column mapping."""
import csv
import io
from dataclasses import dataclass, field

import xlrd
from openpyxl import load_workbook


@dataclass
class SheetRow:
    """A single spreadsheet row of raw cell values."""
    values: list[object] = field(default_factory=list)

    def __len__(self) -> int:
        return len(self.values)

    def __getitem__(self, idx: int) -> object:
        return self.values[idx] if 0 <= idx < len(self.values) else None


class ExcelReader:
    """Strategy-style reader: detects the workbook format and returns rows."""

    def read(self, file_bytes: bytes, filename: str) -> list[SheetRow]:
        name = (filename or "").lower()
        if name.endswith(".xlsx") or name.endswith(".xlsm"):
            return self._read_xlsx(io.BytesIO(file_bytes))
        if name.endswith(".xls"):
            return self._read_xls(io.BytesIO(file_bytes))
        raise ValueError(f"Unsupported file type: {filename}")

    def read_pasted(self, text: str) -> list[SheetRow]:
        """Parses clipboard content (TSV, CSV, or semicolon CSV) into rows."""
        rows: list[SheetRow] = []
        sample = text[:4096]
        delimiter = "\t"
        if "\t" not in sample:
            for candidate in (",", ";", "|"):
                if sample.count(candidate) > sample.count("\n"):
                    delimiter = candidate
                    break
        reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
        for row in reader:
            rows.append(SheetRow(list(row)))
        return rows

    # -- readers ------------------------------------------------------------
    def _read_xls(self, stream: io.BytesIO) -> list[SheetRow]:
        workbook = xlrd.open_workbook(file_contents=stream.read())
        sheet = workbook.sheet_by_index(0)
        rows: list[SheetRow] = []
        for r in range(sheet.nrows):
            rows.append(SheetRow([sheet.cell_value(r, c) for c in range(sheet.ncols)]))
        return rows

    def _read_xlsx(self, stream: io.BytesIO) -> list[SheetRow]:
        # read_only=True truncates some real-world xlsx exports to a single row,
        # so we parse the full workbook instead.
        workbook = load_workbook(stream, data_only=True)
        sheet = workbook.worksheets[0]
        rows: list[SheetRow] = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(SheetRow([v for v in row]))
        workbook.close()
        return rows
